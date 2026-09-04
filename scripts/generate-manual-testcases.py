#!/usr/bin/env python3
"""
Generate manual API test-case Excel workbooks from the TestNG automation suite.

Output:
  docs/qa/manual-testcases/API-Manual-TestCases.xlsx  (one sheet per module + index + duplicacy)
  docs/qa/manual-testcases/modules/*.csv              (per-module CSV for easy sharing)
"""

from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
TESTS_DIR = ROOT / "src" / "test" / "java" / "com" / "automation" / "api" / "tests"
OUT_DIR = ROOT / "docs" / "qa" / "manual-testcases"
OUT_XLSX = OUT_DIR / "API-Manual-TestCases.xlsx"
MODULES_DIR = OUT_DIR / "modules"

HEADERS = [
    "Test Case ID",
    "Module",
    "Feature",
    "API Endpoint",
    "Scenario",
    "Preconditions",
    "Test Data",
    "Test Steps",
    "Expected Result (API)",
    "Expected Result (UI)",
    "Priority",
    "Test Type",
    "Automation Status",
    "Automation Reference",
    "Remarks",
]

# Known data-provider expansions (module -> provider name -> list of variant labels)
DATA_PROVIDER_VARIANTS: dict[str, dict[str, list[str]]] = {
    "ContentDetail.java": {
        "seriesEpisodeViewAllTypes": ["NEXT", "PREVIOUS", "PREVIOUS_NEXT"],
        "bingeWatchEpisodeDirections": ["NEXT", "PREVIOUS"],
        "episodeHierarchyThreeStateDirections": ["NEXT", "CURRENT", "PREVIOUS"],
        "trailerContent": [
            "movie (movie id)",
            "series (close series id)",
            "series (open series id)",
            "boxset (boxset id)",
        ],
        "channelNeighborStates": ["NEXT", "PREVIOUS"],
    },
    "ConfigService.java": {
        "platformConfigPlatforms": [
            "SET_TOP_BOX_ULTRA",
            "SET_TOP_BOX",
            "SET_TOP_BOX_ULTI",
            "ANDROID",
            "IOS",
            "TABLET",
            "IPAD",
            "WEB",
            "ANDROID_TV",
            "LG_HTML_TV",
            "SAMSUNG_HTML_TV",
        ],
        "avatarPagination": ["limit=15, offset=0", "limit=10, offset=0", "limit=20, offset=0"],
    },
    "Footer.java": {
        "footerPlatforms": [
            "STB - Ultra V1",
            "STB - Ultra V2",
            "STB - Ulti Box",
            "Web",
            "Android",
            "iOS",
            "TV",
        ],
    },
    "RecommendationProxy.java": {
        "railUsecases": [
            "dont_miss",
            "popular_top_10",
            "trending_sports",
            "personalised_boxset",
            "popular_top_10_kids",
            "trending_shows",
            "trending_movie_kids",
            "featured_hero",
            "recently_added",
            "Because_You_Watched",
            "top_10_home",
        ],
    },
    "VRSearchProxy.java": {
        "searchUsecases": [
            "trending_on_astro",
            "popular_search",
            "recommend_subgenre_comedy",
            "recommend_subgenre_action",
            "top_10_home",
        ],
    },
}

MODULE_DISPLAY_NAMES = {
    "ContentDetail": "Content Detail",
    "ContinueWatch": "Continue Watch",
    "ContinueWatchScenarios": "Continue Watch Scenarios",
    "WatchAgain": "Watch Again",
    "WatchlistFavourite": "Watchlist & Favourites",
    "BoxsetBingeWatchScenarios": "Boxset Binge Watch",
    "LearnAction": "Learn Action",
    "RecommendationProxy": "Recommendation Proxy",
    "VRSearchProxy": "VR Search Proxy",
    "LockChannel": "Lock Channel",
    "LastTunedChannel": "Last Tuned Channel",
    "SearchHistory": "Search History",
    "HomescreenMenu": "Homescreen Menu",
    "HomescreenProxy": "Homescreen Proxy",
    "ConfigService": "Config Service",
    "Footer": "Footer",
    "TokenGenerator": "Token Generator",
}

MODULE_ORDER = [
    "Content Detail",
    "Continue Watch",
    "Continue Watch Scenarios",
    "Watch Again",
    "Watchlist & Favourites",
    "Boxset Binge Watch",
    "Learn Action",
    "Recommendation Proxy",
    "VR Search Proxy",
    "Lock Channel",
    "Last Tuned Channel",
    "Search History",
    "Homescreen Menu",
    "Homescreen Proxy",
    "Config Service",
    "Footer",
    "Token Generator",
]


@dataclass
class TestCase:
    module: str
    feature: str
    api_endpoint: str
    scenario: str
    method_name: str
    source_file: str
    variant: str = ""
    priority: str = "P1"
    test_type: str = "Positive"
    preconditions: str = ""
    test_data: str = ""
    steps: str = ""
    expected_api: str = ""
    expected_ui: str = ""
    automation_status: str = "Automated"
    remarks: str = ""

    @property
    def dedup_key(self) -> str:
        ep = normalize(self.api_endpoint)
        sc = normalize(self.scenario)
        var = normalize(self.variant)
        return f"{ep}|{sc}|{var}"

    def test_case_id(self, seq: int) -> str:
        prefix = "".join(w[0].upper() for w in self.module.split() if w)[:6]
        return f"{prefix}-{seq:03d}"


def normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def infer_test_type(scenario: str, method_name: str) -> str:
    text = f"{scenario} {method_name}".lower()
    if any(k in text for k in ("400", "negative", "expect http 400", "absent", "not appear", "exclude")):
        return "Negative"
    if any(k in text for k in ("boundary", "96.99", "97%", "threshold")):
        return "Boundary"
    if "teardown" in text or "prepare" in text or "clean slate" in text or "clear" in text:
        return "Setup/Teardown"
    return "Positive"


def infer_priority(scenario: str, test_type: str) -> str:
    if test_type == "Setup/Teardown":
        return "P3"
    if test_type == "Negative":
        return "P2"
    return "P1"


def build_steps(tc: TestCase) -> str:
    if tc.steps:
        return tc.steps
    parts = []
    if tc.api_endpoint:
        parts.append(f"1. Call API: {tc.api_endpoint}")
    if tc.variant:
        parts.append(f"2. Use variant/parameters: {tc.variant}")
    parts.append(f"{len(parts) + 1}. Validate HTTP status, response envelope (status=true), and business rules per scenario.")
    return "\n".join(parts)


def build_expected_api(tc: TestCase) -> str:
    if tc.expected_api:
        return tc.expected_api
    if tc.test_type == "Negative":
        return "API returns expected error (e.g. HTTP 400) or item excluded from list per contract."
    if tc.test_type == "Setup/Teardown":
        return "Setup/teardown completes successfully; state ready for next step."
    return "HTTP 200; status=true; response data matches contract (fields present, correct ids, ordering, no duplicates)."


def parse_java_test_file(path: Path) -> list[TestCase]:
    content = path.read_text(encoding="utf-8")
    class_name = path.stem
    module = MODULE_DISPLAY_NAMES.get(class_name, class_name)

    feature_match = re.search(r'@Feature\("([^"]+)"\)', content)
    feature = feature_match.group(1) if feature_match else module

    cases: list[TestCase] = []
    # Match @Test block then following @Story and method name
    pattern = re.compile(
        r'@Test\s*\((.*?)\)\s*(?:@\w+\([^)]*\)\s*)*@Story\("([^"]+)"\)\s*'
        r'(?:public\s+)?(?:void\s+)?(\w+)\s*\(',
        re.DOTALL,
    )
    alt_pattern = re.compile(
        r'@Test\s*\((.*?)\)\s*public\s+(?:void\s+)?(\w+)\s*\(',
        re.DOTALL,
    )

    for m in pattern.finditer(content):
        test_attrs, story, method = m.group(1), m.group(2), m.group(3)
        desc = extract_description(test_attrs)
        scenario = desc or method.replace("_", " ")
        dp = extract_data_provider(test_attrs)
        variants = expand_variants(path.name, dp)
        if not variants:
            variants = [""]
        for var in variants:
            scenario_full = f"{scenario}" + (f" [{var}]" if var else "")
            tt = infer_test_type(scenario_full, method)
            tc = TestCase(
                module=module,
                feature=feature,
                api_endpoint=story,
                scenario=scenario_full,
                method_name=method,
                source_file=path.name,
                variant=var,
                test_type=tt,
                priority=infer_priority(scenario_full, tt),
                preconditions=default_preconditions(module, scenario_full),
                test_data=default_test_data(module, scenario_full, var),
            )
            cases.append(tc)

    # Tests without @Story (e.g. some with Story after Test on separate pattern)
    for m in alt_pattern.finditer(content):
        method = m.group(2)
        if any(c.method_name == method for c in cases):
            continue
        test_attrs = m.group(1)
        desc = extract_description(test_attrs)
        scenario = desc or method.replace("_", " ")
        # Try to find Story near method
        story_pat = re.search(
            rf'@Story\("([^"]+)"\)\s*(?:public\s+)?(?:void\s+)?{re.escape(method)}\s*\(',
            content,
        )
        story = story_pat.group(1) if story_pat else ""
        dp = extract_data_provider(test_attrs)
        variants = expand_variants(path.name, dp)
        if not variants:
            variants = [""]
        for var in variants:
            scenario_full = f"{scenario}" + (f" [{var}]" if var else "")
            tt = infer_test_type(scenario_full, method)
            cases.append(
                TestCase(
                    module=module,
                    feature=feature,
                    api_endpoint=story,
                    scenario=scenario_full,
                    method_name=method,
                    source_file=path.name,
                    variant=var,
                    test_type=tt,
                    priority=infer_priority(scenario_full, tt),
                    preconditions=default_preconditions(module, scenario_full),
                    test_data=default_test_data(module, scenario_full, var),
                )
            )

    return cases


def extract_description(attrs: str) -> str:
    m = re.search(r'description\s*=\s*"([^"]+)"', attrs, re.DOTALL)
    if m:
        return re.sub(r"\s+", " ", m.group(1).strip())
    return ""


def extract_data_provider(attrs: str) -> str | None:
    m = re.search(r'dataProvider\s*=\s*"([^"]+)"', attrs)
    return m.group(1) if m else None


def expand_variants(filename: str, provider: str | None) -> list[str]:
    if not provider:
        return []
    file_map = DATA_PROVIDER_VARIANTS.get(filename, {})
    return file_map.get(provider, [f"dataProvider={provider}"])


def default_preconditions(module: str, scenario: str) -> str:
    s = scenario.lower()
    if "guest" in s:
        return "Guest user session with valid guest bearer token and x-api-key."
    if "teardown" in s or "prepare" in s or "clean slate" in s or "clear" in s:
        return "Logged-in subscriber with valid auth headers."
    if module in ("Continue Watch", "Continue Watch Scenarios", "Watch Again", "Boxset Binge Watch"):
        return "Logged-in entitled subscriber; CW/WA lists in known state (see scenario)."
    if module == "Lock Channel":
        return "Logged-in subscriber; parental lock PIN configured if required by client."
    return "Logged-in subscriber with valid bearer token, x-api-key, and environment-specific test content configured."


def default_test_data(module: str, scenario: str, variant: str) -> str:
    if variant:
        return f"Environment properties + variant: {variant}"
    if "channel" in scenario.lower():
        return "vrgo.lock.channel.* / vrgo.last.tuned.channel.* from environment properties"
    if module == "Content Detail":
        return "vrgo.content.detail.* ids from environments/<env>.properties"
    if module in ("Recommendation Proxy", "VR Search Proxy"):
        return "vrgo.recommendation.proxy.* / vrgo.search.proxy.* from environment properties"
    return "Configured content/channel ids from environments/<env>.properties"


def load_manual_only_boxset_cases() -> list[TestCase]:
    """Manual-only cases documented in docs/qa/boxset-binge-watch-cw-wa-test-suite.md."""
    manual_rows = [
        ("BB-CW-M03", "Shared movie M1 replay after completion", "M1 in WA; siblings in CW", "M1", "10%", "false", "P2", "Regression"),
        ("BB-CW-M04", "Remove M1 playback history", "M1 completed", "M1", "—", "—", "P2", "Negative"),
        ("BB-CW-T03", "Resume T1 after completion", "T1 in WA", "T1", "5%", "false", "P2", "Regression"),
        ("BB-CW-T04", "Replay completed T1", "T1 in WA", "T1", "100%", "true", "P2", "Regression"),
        ("BB-BND-01", "M1 at 96.99% — below threshold", "Empty lists", "M1", "96.99%", "false", "P1", "Boundary"),
        ("BB-BND-02", "M1 at 97.00% — at threshold", "Empty lists", "M1", "97.00%", "false", "P1", "Boundary"),
        ("BB-BND-03", "M1 at 97.01% — above threshold", "Empty lists", "M1", "97.01%", "false", "P1", "Boundary"),
        ("BB-CLF-01", "M1: 96.99% + hasCompletedPlayback=true", "Empty lists", "M1", "96.99%", "true", "P1", "Edge"),
        ("BB-REG-05", "Refresh application — CW/WA persistence", "M1 complete state", "M1", "—", "—", "P1", "Regression"),
        ("BB-REG-06", "Logout/Login — CW/WA restored", "M1 complete state", "M1", "—", "—", "P1", "Regression"),
        ("BB-UI-01", "Correct cards on CW rail (UI)", "M1 partial in CW", "M1", "50%", "false", "P1", "UI"),
        ("BB-UI-02", "Correct cards on WA rail (UI)", "M1 complete", "M1", "100%", "true", "P1", "UI"),
    ]
    cases = []
    for row in manual_rows:
        tc_id, scenario, pre, data, pct, flag, pri, ttype = row
        cases.append(
            TestCase(
                module="Boxset Binge Watch",
                feature="Boxset Binge CW/WA",
                api_endpoint="Manual / UI + CW/WA APIs",
                scenario=f"{scenario} ({tc_id})",
                method_name=tc_id,
                source_file="boxset-binge-watch-cw-wa-test-suite.md",
                priority=pri,
                test_type=ttype,
                preconditions=pre,
                test_data=f"{data}; playback={pct}; hasCompletedPlayback={flag}",
                automation_status="Manual Only",
                remarks="From QA doc; not fully automated in BoxsetBingeWatchScenarios",
            )
        )
    return cases


def to_row(tc: TestCase, seq: int) -> list[str]:
    return [
        tc.test_case_id(seq),
        tc.module,
        tc.feature,
        tc.api_endpoint,
        tc.scenario,
        tc.preconditions,
        tc.test_data,
        build_steps(tc),
        build_expected_api(tc),
        tc.expected_ui or ("Verify client UI matches API state where applicable." if tc.test_type == "UI" else "N/A (API contract test)"),
        tc.priority,
        tc.test_type,
        tc.automation_status,
        f"{tc.source_file}#{tc.method_name}",
        tc.remarks,
    ]


def analyze_duplicates(all_cases: list[TestCase]) -> list[list[str]]:
    """Return rows for duplicacy sheet."""
    by_key: dict[str, list[TestCase]] = defaultdict(list)
    by_endpoint: dict[str, list[TestCase]] = defaultdict(list)
    by_scenario: dict[str, list[TestCase]] = defaultdict(list)

    for tc in all_cases:
        if tc.test_type == "Setup/Teardown":
            continue
        by_key[tc.dedup_key].append(tc)
        ep = normalize(tc.api_endpoint)
        if ep:
            by_endpoint[ep].append(tc)
        by_scenario[normalize(tc.scenario)].append(tc)

    rows: list[list[str]] = []
    rows.append([
        "Duplicate Group",
        "Severity",
        "Modules Involved",
        "API Endpoint",
        "Scenario",
        "Count",
        "Automation References",
        "Recommendation",
    ])

    group_num = 0
    seen_groups: set[frozenset[str]] = set()

    # Exact duplicate: same endpoint + scenario + variant across modules
    for key, group in sorted(by_key.items()):
        if len(group) < 2:
            continue
        modules = sorted({tc.module for tc in group})
        if len(modules) < 2 and len(group) < 2:
            continue
        refs = sorted({f"{tc.module}:{tc.method_name}" for tc in group})
        group_id = frozenset(refs)
        if group_id in seen_groups:
            continue
        seen_groups.add(group_id)
        group_num += 1
        rows.append([
            f"DUP-{group_num:03d}",
            "High" if len(modules) > 1 else "Medium",
            ", ".join(modules),
            group[0].api_endpoint,
            group[0].scenario,
            str(len(group)),
            "; ".join(refs),
            "Consolidate into one test case; keep single automation + one manual row." if len(modules) > 1
            else "Review if multiple methods test same behavior; merge or differentiate expected results.",
        ])

    # Cross-module endpoint overlap (same API, different scenarios)
    cross_module_endpoints = []
    for ep, group in by_endpoint.items():
        modules = {tc.module for tc in group}
        if len(modules) > 1 and ep:
            cross_module_endpoints.append((ep, sorted(modules), len(group)))

    for ep, modules, count in sorted(cross_module_endpoints, key=lambda x: -x[2])[:25]:
        group_num += 1
        rows.append([
            f"DUP-{group_num:03d}",
            "Info",
            ", ".join(modules),
            ep,
            "(multiple scenarios)",
            str(count),
            "",
            "Shared API across modules — expected for CW/WA flows; ensure manual sheet references master module.",
        ])

    # Known semantic duplicates called out explicitly
    known = [
        (
            "GET boxset binge API",
            ["Content Detail", "Boxset Binge Watch"],
            "GET /content-detail-service/pub/v1/boxset/{boxsetId}/binge",
            "Same binge API validated in ContentDetail and BoxsetBingeWatchScenarios — keep Content Detail as contract test; Boxset as E2E.",
        ),
        (
            "CW list GET",
            ["Continue Watch", "Continue Watch Scenarios", "Watch Again", "Boxset Binge Watch"],
            "GET /subscriber-event-service/v3/continue-watch/continue",
            "Multiple modules GET same CW list for different flows — not duplicate tests if scenarios differ; tag with flow name in manual sheet.",
        ),
        (
            "WA list GET",
            ["Watch Again", "Continue Watch Scenarios", "Boxset Binge Watch"],
            "GET /subscriber-event-service/v3/watch-again",
            "Consolidate manual expected results per scenario; avoid copy-paste expected CW/WA lists.",
        ),
        (
            "VR Search usecase duplicate",
            ["VR Search Proxy"],
            "GET search-by-usecase",
            "searchUsecases array lists recommend_subgenre_comedy twice in automation — remove duplicate in code or manual sheet.",
        ),
    ]
    for title, modules, endpoint, rec in known:
        group_num += 1
        rows.append([
            f"DUP-{group_num:03d}",
            "Review",
            ", ".join(modules),
            endpoint,
            title,
            "",
            "",
            rec,
        ])

    return rows


def style_sheet(ws, header_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, header_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(header_count)}1"
    for col in range(1, header_count + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["H"].width = 35
    ws.column_dimensions["I"].width = 35


def write_csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        writer.writerows(rows)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    MODULES_DIR.mkdir(parents=True, exist_ok=True)

    all_cases: list[TestCase] = []
    for java_file in sorted(TESTS_DIR.glob("*.java")):
        if java_file.name == "TokenGenerator.java":
            continue  # infra; optional in manual suite
        all_cases.extend(parse_java_test_file(java_file))

    all_cases.extend(load_manual_only_boxset_cases())

    by_module: dict[str, list[TestCase]] = defaultdict(list)
    for tc in all_cases:
        by_module[tc.module].append(tc)

    wb = Workbook()
    wb.remove(wb.active)

    # Index sheet
    idx = wb.create_sheet("Index", 0)
    idx.append(["Module", "Sheet Name", "Test Case Count", "Automated", "Manual Only"])
    for module in MODULE_ORDER:
        if module not in by_module:
            continue
        cases = by_module[module]
        auto = sum(1 for c in cases if c.automation_status == "Automated")
        manual = len(cases) - auto
        sheet_name = module[:31]
        idx.append([module, sheet_name, len(cases), auto, manual])
    idx.append([])
    idx.append(["Total", "", len(all_cases), "", ""])
    style_sheet(idx, 5)

    # Module sheets + CSV
    for module in MODULE_ORDER:
        if module not in by_module:
            continue
        cases = by_module[module]
        sheet_name = module[:31]
        ws = wb.create_sheet(sheet_name)
        ws.append(HEADERS)
        rows = []
        for i, tc in enumerate(cases, start=1):
            row = to_row(tc, i)
            ws.append(row)
            rows.append(row)
        style_sheet(ws, len(HEADERS))
        csv_name = re.sub(r"[^\w\-]+", "_", module).strip("_") + ".csv"
        write_csv(MODULES_DIR / csv_name, rows)

    # Duplicacy sheet
    dup_ws = wb.create_sheet("Duplicacy Analysis")
    for row in analyze_duplicates(all_cases):
        dup_ws.append(row)
    style_sheet(dup_ws, 8)
    dup_ws.column_dimensions["G"].width = 45
    dup_ws.column_dimensions["H"].width = 50

    wb.save(OUT_XLSX)
    print(f"Generated: {OUT_XLSX}")
    print(f"Module CSVs: {MODULES_DIR}")
    print(f"Total test cases: {len(all_cases)}")
    print(f"Modules: {len(by_module)}")


if __name__ == "__main__":
    main()
